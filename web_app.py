"""Servidor web para probar el proyecto de aprendizaje supervisado.

Este módulo sirve una interfaz gráfica desde la carpeta web/ y expone rutas
para generar un dataset sintético, entrenar modelos, consultar métricas y
exportar resultados como un libro de Excel desde el navegador.
"""

from __future__ import annotations

import json
from io import BytesIO
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT_DIR = Path(__file__).resolve().parent
WEB_DIR = ROOT_DIR / "web"
APP_DIR = ROOT_DIR / "aprendizaje_supervisado"

# Se incorpora la raíz del proyecto para poder importar los módulos de Python.
sys.path.insert(0, str(ROOT_DIR))

from aprendizaje_supervisado.generate_dataset import generate_dataset
from aprendizaje_supervisado.train_model import train_and_evaluate


class ProjectHandler(BaseHTTPRequestHandler):
    """Maneja las peticiones HTTP del navegador y sirve la interfaz gráfica."""

    def do_GET(self) -> None:
        """Responde con la página principal o archivos estáticos del frontend."""
        parsed = urlparse(self.path)
        route = parsed.path

        if route in ("/", "/index.html"):
            self._serve_file(WEB_DIR / "index.html", "text/html; charset=utf-8")
            return

        if route == "/api/metrics":
            self._send_json(self._read_metrics())
            return

        if route.startswith("/"):
            relative_path = route.lstrip("/")
            file_path = WEB_DIR / relative_path
            if file_path.exists() and file_path.is_file():
                content_type = self._content_type_for(file_path)
                self._serve_file(file_path, content_type)
                return

        self._send_json({"error": "Ruta no encontrada."}, status=404)

    def do_POST(self) -> None:
        """Procesa generación de datos, entrenamiento y exportación a Excel."""
        parsed = urlparse(self.path)
        route = parsed.path

        if route == "/api/generate-dataset":
            payload = self._read_json_body()
            sample_size = int(payload.get("sample_size", 1200))
            seed = int(payload.get("seed", 2026))
            dataset_path = generate_dataset(sample_size=sample_size, seed=seed)
            self._send_json(
                {
                    "status": "ok",
                    "message": "Dataset generado correctamente.",
                    "sample_size": sample_size,
                    "seed": seed,
                    "path": str(dataset_path.relative_to(ROOT_DIR)),
                }
            )
            return

        if route == "/api/train-model":
            metrics = train_and_evaluate()
            self._send_json(
                {
                    "status": "ok",
                    "message": "Entrenamiento finalizado.",
                    "results": metrics,
                    "metrics_path": str((APP_DIR / "artifacts" / "metrics.json").relative_to(ROOT_DIR)),
                }
            )
            return

        if route == "/api/export-comparison":
            payload = self._read_json_body()
            # Acepta tanto la respuesta del entrenamiento como el reporte persistido.
            metrics = payload.get("results") or payload.get("models") or payload
            workbook = self._create_comparison_workbook(metrics)
            if workbook is None:
                self._send_json(
                    {"error": "No hay métricas válidas para exportar."}, status=400
                )
                return
            self._send_file_content(
                workbook,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "comparacion_modelos.xlsx",
            )
            return

        self._send_json({"error": "Endpoint no disponible."}, status=404)

    def log_message(self, format: str, *args: object) -> None:
        """Suprime el ruido habitual de los logs de la librería HTTP."""
        return

    def _read_json_body(self) -> dict:
        """Lee y parsea el contenido JSON enviado por el navegador."""
        content_length = int(self.headers.get("Content-Length", "0"))
        raw_data = self.rfile.read(content_length)
        if not raw_data:
            return {}
        return json.loads(raw_data.decode("utf-8"))

    def _read_metrics(self) -> dict:
        """Carga el archivo de métricas generado por el entrenamiento."""
        metrics_path = APP_DIR / "artifacts" / "metrics.json"
        if not metrics_path.exists():
            return {"status": "no-data", "message": "Todavía no hay métricas disponibles."}
        return json.loads(metrics_path.read_text(encoding="utf-8"))

    def _serve_file(self, file_path: Path, content_type: str) -> None:
        """Sirve un archivo del disco como respuesta HTTP."""
        content = file_path.read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(content)))
        self.end_headers()
        self.wfile.write(content)

    def _send_json(self, payload: dict, status: int = 200) -> None:
        """Devuelve una respuesta JSON con el código apropiado."""
        body = json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _send_file_content(self, content: bytes, content_type: str, filename: str) -> None:
        """Envía contenido binario al navegador como un archivo descargable."""
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
        self.send_header("Content-Length", str(len(content)))
        self.end_headers()
        self.wfile.write(content)

    @staticmethod
    def _create_comparison_workbook(metrics: dict) -> bytes | None:
        """Crea un libro Excel con las métricas y resalta el mejor MAE."""
        # Traduce los nombres internos de los modelos a etiquetas legibles en Excel.
        model_names = {
            "baseline_median": "Baseline",
            "linear_regression": "Regresión lineal",
            "random_forest": "Bosque aleatorio",
        }
        valid_models = {
            name: metrics[name]
            for name in model_names
            if isinstance(metrics, dict) and isinstance(metrics.get(name), dict)
        }
        if not valid_models:
            return None

        # Crea una hoja con encabezados destacados y una fila para cada modelo disponible.
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Comparación"
        sheet.append(["Modelo", "MAE (minutos)", "RMSE (minutos)", "R²"])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="245C58")

        for key, label in model_names.items():
            result = valid_models.get(key)
            if result is None:
                continue
            sheet.append(
                [label, result.get("mae_minutes"), result.get("rmse_minutes"), result.get("r2")]
            )

        # Resalta el menor error absoluto medio, que es la métrica principal de comparación.
        mae_rows = [
            (row, sheet.cell(row=row, column=2).value)
            for row in range(2, sheet.max_row + 1)
            if isinstance(sheet.cell(row=row, column=2).value, (int, float))
        ]
        if mae_rows:
            best_row = min(mae_rows, key=lambda item: item[1])[0]
            for cell in sheet[best_row]:
                cell.fill = PatternFill("solid", fgColor="D9EAD3")

        for column, width in {"A": 24, "B": 18, "C": 20, "D": 12}.items():
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A2"

        # Serializa el libro en memoria para enviarlo directamente al navegador.
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    @staticmethod
    def _content_type_for(file_path: Path) -> str:
        """Determina el tipo MIME real según la extensión del archivo."""
        mime_type, _ = __import__("mimetypes").guess_type(str(file_path))
        if mime_type is None:
            return "application/octet-stream"
        return mime_type


if __name__ == "__main__":
    host = "127.0.0.1"
    # Se usa un puerto libre porque el 8000 aún puede tener una instancia antigua activa.
    port = 8001
    server = ThreadingHTTPServer((host, port), ProjectHandler)
    print(f"Servidor web activo en http://{host}:{port}")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nServidor detenido.")
        server.server_close()
